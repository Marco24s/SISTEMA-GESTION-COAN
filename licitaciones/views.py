import csv
import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.contrib.auth.mixins import LoginRequiredMixin, UserPassesTestMixin
from django.contrib.messages.views import SuccessMessageMixin
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render

class GroupRequiredMixin(UserPassesTestMixin):
    group_required = None

    def test_func(self):
        if self.request.user.is_superuser:
            return True
        if not self.group_required:
            return True
        groups = self.group_required if isinstance(self.group_required, list) else [self.group_required]
        return self.request.user.groups.filter(name__in=groups).exists()
from django.db.models import Count, Q, Sum
from django.shortcuts import get_object_or_404
from django.urls import reverse_lazy
from django.utils import timezone
from django.views.generic import CreateView, DetailView, ListView, TemplateView, UpdateView
from django.views import View

from core.models import Unit

from .forms import (
    ForeignTenderProcessForm,
    ForeignTenderPurchaseOrderForm,
    ForeignTenderRequirementForm,
    ForeignTenderUpdateForm,
    ForeignProvisionRequestForm,
    TenderProcessForm,
    TenderStageUpdateForm,
)
from .models import (
    ForeignTenderProcess,
    ForeignTenderPurchaseOrder,
    ForeignTenderRequirement,
    ForeignTenderUpdate,
    ForeignProvisionRequest,
    ProcurementDestination,
    TenderProcess,
    TenderStage,
)


class TenderTypeSelectionView(LoginRequiredMixin, TemplateView):
    template_name = "licitaciones/type_selection.html"


def _clean_int(value):
    if value in (None, ""):
        return None
    text = str(value).strip().replace(".", "").replace(",", "")
    return int(text) if text.isdigit() else None


def _percent(part, total):
    if not total:
        return 0
    return round((part * 100) / total, 1)


def _status_filter_for_group(group):
    if group == "ADJUDICADO":
        return ["ADJUDICADO"]
    if group == "DISPONIBLE":
        return ["PREADJUDICADO", "DISPONIBLE_ADJUDICAR", "PREADJUDICADO_DISPONIBLE"]
    if group == "EN_PROCESO":
        return ["PUBLICADO", "EN_APERTURA", "EN_EVALUACION"]
    if group == "SIN_EFECTO":
        return ["FRACASADO", "DESIERTO", "DEJADO_SIN_EFECTO"]
    return []


def _get_classification_choices():
    default_keys = [c[0] for c in TenderProcess.CLASSIFICATION_CHOICES]
    choices = [choice for choice in TenderProcess.CLASSIFICATION_CHOICES if choice[0]]
    existing_custom = (
        TenderProcess.objects.exclude(classification__in=default_keys)
        .exclude(classification__isnull=True)
        .exclude(classification="")
        .values_list("classification", flat=True)
        .distinct()
        .order_by("classification")
    )
    for c in existing_custom:
        if (c, c) not in choices:
            choices.append((c, c))
    return choices


def _get_custom_classifications_with_count():
    default_keys = ["", "REPUESTO", "SUPERVIVENCIA", "GRASAS_LUBRICANTES", "REPUESTOS_FONDEF"]
    items = (
        TenderProcess.objects.exclude(classification__in=default_keys)
        .exclude(classification__isnull=True)
        .exclude(classification="")
        .values("classification")
        .annotate(count=Count("id"))
        .order_by("classification")
    )
    return [{"name": item["classification"], "count": item["count"]} for item in items]


class TenderDashboardView(LoginRequiredMixin, TemplateView):
    template_name = "licitaciones/dashboard.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        years = list(
            TenderProcess.objects.filter(is_active=True)
            .order_by("-year")
            .values_list("year", flat=True)
            .distinct()
        )
        selected_year = _clean_int(self.request.GET.get("year"))
        if not selected_year and years:
            selected_year = years[0]

        processes = TenderProcess.objects.select_related("unit", "destination").filter(is_active=True)
        if selected_year:
            processes = processes.filter(year=selected_year)

        process_list = list(processes)
        total_count = len(process_list)
        adjudicated = [p for p in process_list if p.operational_group == "ADJUDICADO"]
        available = [p for p in process_list if p.operational_group == "DISPONIBLE"]
        cancelled = [p for p in process_list if p.operational_group == "SIN_EFECTO"]
        in_progress = [p for p in process_list if p.operational_group == "EN_PROCESO"]

        def amount_sum(items):
            return sum((p.amount_ars or 0) for p in items)

        destination_rows = []
        units = Unit.objects.filter().order_by("name")
        for unit in units:
            items = [p for p in process_list if p.unit_id == unit.id]
            item_count = len(items)
            if item_count == 0:
                continue
            destination_rows.append(
                {
                    "unit": unit,
                    "total": item_count,
                    "adjudicated_count": len([p for p in items if p.operational_group == "ADJUDICADO"]),
                    "adjudicated_amount": amount_sum([p for p in items if p.operational_group == "ADJUDICADO"]),
                    "available_count": len([p for p in items if p.operational_group == "DISPONIBLE"]),
                    "available_amount": amount_sum([p for p in items if p.operational_group == "DISPONIBLE"]),
                    "in_progress_count": len([p for p in items if p.operational_group == "EN_PROCESO"]),
                    "cancelled_count": len([p for p in items if p.operational_group == "SIN_EFECTO"]),
                    "percent": _percent(item_count, total_count),
                }
            )

        status_rows = []
        for value, label in TenderProcess.STATUS_CHOICES:
            items = [p for p in process_list if p.status == value]
            count = len(items)
            if count:
                status_rows.append(
                    {
                        "value": value,
                        "label": label,
                        "count": count,
                        "amount": amount_sum(items),
                        "percent": _percent(count, total_count),
                    }
                )

        today = timezone.localdate()
        upcoming_opening_alerts = []
        overdue_opening_alerts = []
        opening_processes = [
            p
            for p in process_list
            if p.opening_date and p.operational_group in ["EN_PROCESO", "DISPONIBLE"]
        ]
        for process in sorted(opening_processes, key=lambda p: p.opening_date):
            opening_date = timezone.localtime(process.opening_date).date()
            days_until = (opening_date - today).days
            if days_until > 7:
                break
            if days_until < 0:
                days_late = abs(days_until)
                alert_label = f"Vencida hace {days_late} dia" if days_late == 1 else f"Vencida hace {days_late} dias"
                alert_class = "text-bg-danger"
            elif days_until == 0:
                alert_label = "Abre hoy"
                alert_class = "text-bg-warning"
            elif days_until <= 7:
                alert_label = f"En {days_until} dia" if days_until == 1 else f"En {days_until} dias"
                alert_class = "text-bg-info"
            else:
                alert_label = f"En {days_until} dia(s)"
                alert_class = "text-bg-light border"
            alert = {
                "process": process,
                "alert_label": alert_label,
                "alert_class": alert_class,
                "days_until": days_until,
            }
            if days_until < 0:
                overdue_opening_alerts.append(alert)
            else:
                upcoming_opening_alerts.append(alert)

        overdue_opening_alerts = sorted(overdue_opening_alerts, key=lambda item: item["days_until"], reverse=True)
        opening_alerts = upcoming_opening_alerts + overdue_opening_alerts[:5]

        context.update(
            {
                "years": years,
                "selected_year": selected_year or "",
                "total_count": total_count,
                "total_amount": amount_sum(process_list),
                "adjudicated_count": len(adjudicated),
                "adjudicated_amount": amount_sum(adjudicated),
                "adjudicated_percent": _percent(len(adjudicated), total_count),
                "available_count": len(available),
                "available_amount": amount_sum(available),
                "available_percent": _percent(len(available), total_count),
                "in_progress_count": len(in_progress),
                "in_progress_percent": _percent(len(in_progress), total_count),
                "cancelled_count": len(cancelled),
                "cancelled_percent": _percent(len(cancelled), total_count),
                "foreign_count": len([p for p in process_list if p.currency in ["USD", "EUR", "OTRA"]]),
                "missing_amount_count": len([p for p in process_list if p.amount_ars is None]),
                "destination_rows": destination_rows,
                "status_rows": status_rows,
                "opening_alerts": opening_alerts,
                "opening_alert_days": 7,
            }
        )
        return context


class TenderProcessListView(LoginRequiredMixin, ListView):
    model = TenderProcess
    template_name = "licitaciones/process_list.html"
    context_object_name = "processes"
    paginate_by = 25

    def get_queryset(self):
        queryset = (
            TenderProcess.objects.select_related("unit", "destination", "created_by")
            .filter(is_active=True)
            .order_by("-year", "unit__name", "-opening_date", "process_number")
        )
        year = _clean_int(self.request.GET.get("year"))
        unit = _clean_int(self.request.GET.get("unit"))
        status = self.request.GET.get("status")
        classification = self.request.GET.get("classification")
        group = self.request.GET.get("group")
        control = self.request.GET.get("control")
        ipp = self.request.GET.get("ipp", "").strip()
        q = self.request.GET.get("q")

        if year:
            queryset = queryset.filter(year=year)
        if unit:
            queryset = queryset.filter(unit_id=unit)
        if status:
            queryset = queryset.filter(status=status)
        elif group:
            group_statuses = _status_filter_for_group(group)
            if group_statuses:
                queryset = queryset.filter(status__in=group_statuses)
        if classification:
            queryset = queryset.filter(classification=classification)
        if control == "missing_amount":
            queryset = queryset.filter(amount_ars__isnull=True)
        elif control == "foreign_currency":
            queryset = queryset.filter(currency__in=["USD", "EUR", "OTRA"])
        if ipp:
            queryset = queryset.filter(ipp__icontains=ipp)
        if q:
            queryset = queryset.filter(
                Q(process_number__icontains=q)
                | Q(expediente__icontains=q)
                | Q(name__icontains=q)
                | Q(ipp__icontains=q)
            )
        return queryset

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context["units"] = Unit.objects.filter().order_by("name")
        context["status_choices"] = TenderProcess.STATUS_CHOICES
        context["classification_choices"] = _get_classification_choices()
        context["selected_year"] = _clean_int(self.request.GET.get("year")) or ""
        context["selected_unit"] = _clean_int(self.request.GET.get("unit")) or ""
        context["selected_status"] = self.request.GET.get("status", "")
        context["selected_classification"] = self.request.GET.get("classification", "")
        context["selected_group"] = self.request.GET.get("group", "")
        context["selected_control"] = self.request.GET.get("control", "")
        context["selected_ipp"] = self.request.GET.get("ipp", "").strip()
        context["search_query"] = self.request.GET.get("q", "")
        context["years"] = (
            TenderProcess.objects.order_by("-year")
            .values_list("year", flat=True)
            .distinct()
        )
        context["ipp_choices"] = (
            TenderProcess.objects.exclude(ipp__isnull=True)
            .exclude(ipp="")
            .values_list("ipp", flat=True)
            .distinct()
            .order_by("ipp")
        )
        context["custom_classifications"] = _get_custom_classifications_with_count()
        query_params = self.request.GET.copy()
        query_params.pop("page", None)
        context["pagination_query"] = query_params.urlencode()
        return context


class TenderProcessHistoryView(LoginRequiredMixin, ListView):
    model = TenderProcess
    template_name = "licitaciones/process_history.html"
    context_object_name = "processes"
    paginate_by = 25

    def get_queryset(self):
        queryset = (
            TenderProcess.objects.select_related("unit", "destination", "created_by")
            .filter(is_active=False)
            .order_by("-year", "unit__name", "-opening_date", "process_number")
        )
        year = _clean_int(self.request.GET.get("year"))
        unit = _clean_int(self.request.GET.get("unit"))
        status = self.request.GET.get("status")
        classification = self.request.GET.get("classification")
        ipp = self.request.GET.get("ipp", "").strip()
        q = self.request.GET.get("q")

        if year:
            queryset = queryset.filter(year=year)
        if unit:
            queryset = queryset.filter(unit_id=unit)
        if status:
            queryset = queryset.filter(status=status)
        if classification:
            queryset = queryset.filter(classification=classification)
        if ipp:
            queryset = queryset.filter(ipp__icontains=ipp)
        if q:
            queryset = queryset.filter(
                Q(process_number__icontains=q)
                | Q(expediente__icontains=q)
                | Q(name__icontains=q)
                | Q(ipp__icontains=q)
            )
        return queryset

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context["units"] = Unit.objects.filter().order_by("name")
        context["status_choices"] = TenderProcess.STATUS_CHOICES
        context["classification_choices"] = _get_classification_choices()
        context["custom_classifications"] = _get_custom_classifications_with_count()
        context["selected_year"] = _clean_int(self.request.GET.get("year")) or ""
        context["selected_unit"] = _clean_int(self.request.GET.get("unit")) or ""
        context["selected_status"] = self.request.GET.get("status", "")
        context["selected_classification"] = self.request.GET.get("classification", "")
        context["selected_ipp"] = self.request.GET.get("ipp", "").strip()
        context["search_query"] = self.request.GET.get("q", "")
        context["years"] = (
            TenderProcess.objects.filter(is_active=False)
            .order_by("-year")
            .values_list("year", flat=True)
            .distinct()
        )
        context["ipp_choices"] = (
            TenderProcess.objects.exclude(ipp__isnull=True)
            .exclude(ipp="")
            .values_list("ipp", flat=True)
            .distinct()
            .order_by("ipp")
        )
        return context


class TenderProcessCreateView(LoginRequiredMixin, GroupRequiredMixin, SuccessMessageMixin, CreateView):
    group_required = ["Supervisor", "Capturista"]
    model = TenderProcess
    form_class = TenderProcessForm
    template_name = "licitaciones/process_form.html"
    success_url = reverse_lazy("licitaciones:process_list")
    success_message = "Proceso licitatorio creado correctamente."

    def form_valid(self, form):
        form.instance.created_by = self.request.user
        return super().form_valid(form)

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context["custom_classifications"] = _get_custom_classifications_with_count()
        return context


class TenderProcessDetailView(LoginRequiredMixin, DetailView):
    model = TenderProcess
    template_name = "licitaciones/process_detail.html"
    context_object_name = "process"

    def get_queryset(self):
        return TenderProcess.objects.select_related("unit", "destination", "created_by")


class TenderProcessUpdateView(LoginRequiredMixin, GroupRequiredMixin, SuccessMessageMixin, UpdateView):
    group_required = ["Supervisor", "Capturista"]
    model = TenderProcess
    form_class = TenderProcessForm
    template_name = "licitaciones/process_form.html"
    success_message = "Proceso licitatorio actualizado correctamente."

    def get_success_url(self):
        if self.object.is_active:
            return reverse_lazy("licitaciones:process_list")
        return reverse_lazy("licitaciones:process_history")

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context["custom_classifications"] = _get_custom_classifications_with_count()
        return context


class TenderClassificationDeleteView(LoginRequiredMixin, GroupRequiredMixin, View):
    group_required = ["Supervisor", "Capturista"]

    def post(self, request, *args, **kwargs):
        classification_name = request.POST.get("classification_name", "").strip()
        if classification_name:
            updated_count = TenderProcess.objects.filter(classification=classification_name).update(classification="")
            messages.success(
                request,
                f"Clasificación '{classification_name}' eliminada correctamente ({updated_count} proceso(s) desvinculado(s))."
            )
        else:
            messages.error(request, "No se especificó ninguna clasificación para eliminar.")

        next_url = request.POST.get("next") or request.META.get("HTTP_REFERER") or reverse("licitaciones:process_list")
        return redirect(next_url)


class ForeignTenderDashboardView(LoginRequiredMixin, TemplateView):
    template_name = "licitaciones/foreign_dashboard.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        processes = ForeignTenderProcess.objects.filter(is_active=True).prefetch_related("requirements", "updates")
        process_list = list(processes)
        totals_by_currency = list(
            ForeignTenderProcess.objects.filter(is_active=True)
            .values("currency", "custom_currency")
            .annotate(
                evaluation_total=Sum("evaluation_amount"),
                awarded_total=Sum("awarded_amount"),
                process_count=Count("id"),
            )
            .order_by("currency")
        )

        context.update(
            {
                "processes": process_list,
                "total_count": len(process_list),
                "active_count": len([p for p in process_list if p.is_active]),
                "awarded_count": len([p for p in process_list if p.status == "ADJUDICADO"]),
                "pending_count": len(
                    [
                        p
                        for p in process_list
                        if p.status
                        not in {"ADJUDICADO", "FINALIZADO", "FRACASADO", "DEJADO_SIN_EFECTO"}
                    ]
                ),
                "totals_by_currency": totals_by_currency,
            }
        )
        return context


class ForeignTenderProcessListView(LoginRequiredMixin, ListView):
    model = ForeignTenderProcess
    template_name = "licitaciones/foreign_list.html"
    context_object_name = "processes"

    def get_queryset(self):
        queryset = ForeignTenderProcess.objects.filter(is_active=True).prefetch_related("requirements", "purchase_orders", "updates")
        year = _clean_int(self.request.GET.get("year"))
        status = self.request.GET.get("status", "")
        currency = self.request.GET.get("currency", "")
        query = self.request.GET.get("q", "").strip()
        if year:
            queryset = queryset.filter(year=year)
        if status:
            queryset = queryset.filter(status=status)
        if currency:
            queryset = queryset.filter(currency=currency)
        if query:
            queryset = queryset.filter(
                Q(process_number__icontains=query)
                | Q(expediente__icontains=query)
                | Q(requirements__requirement_number__icontains=query)
                | Q(requirements__description__icontains=query)
                | Q(purchase_orders__order_number__icontains=query)
                | Q(purchase_orders__supplier__icontains=query)
            ).distinct()
        return queryset

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context.update(
            {
                "years": ForeignTenderProcess.objects.filter(is_active=True).order_by("-year")
                .values_list("year", flat=True)
                .distinct(),
                "status_choices": ForeignTenderProcess.STATUS_CHOICES,
                "currency_choices": ForeignTenderProcess.CURRENCY_CHOICES,
                "selected_year": self.request.GET.get("year", ""),
                "selected_status": self.request.GET.get("status", ""),
                "selected_currency": self.request.GET.get("currency", ""),
                "search_query": self.request.GET.get("q", ""),
            }
        )
        return context


class ForeignTenderProcessHistoryView(LoginRequiredMixin, ListView):
    model = ForeignTenderProcess
    template_name = "licitaciones/foreign_history.html"
    context_object_name = "processes"

    def get_queryset(self):
        queryset = ForeignTenderProcess.objects.filter(is_active=False).prefetch_related(
            "requirements", "purchase_orders", "updates"
        )
        year = _clean_int(self.request.GET.get("year"))
        status = self.request.GET.get("status", "")
        currency = self.request.GET.get("currency", "")
        query = self.request.GET.get("q", "").strip()
        if year:
            queryset = queryset.filter(year=year)
        if status:
            queryset = queryset.filter(status=status)
        if currency:
            queryset = queryset.filter(currency=currency)
        if query:
            queryset = queryset.filter(
                Q(process_number__icontains=query)
                | Q(expediente__icontains=query)
                | Q(requirements__requirement_number__icontains=query)
                | Q(requirements__description__icontains=query)
                | Q(purchase_orders__order_number__icontains=query)
                | Q(purchase_orders__supplier__icontains=query)
            ).distinct()
        return queryset

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context.update(
            {
                "years": ForeignTenderProcess.objects.filter(is_active=False)
                .order_by("-year")
                .values_list("year", flat=True)
                .distinct(),
                "status_choices": ForeignTenderProcess.STATUS_CHOICES,
                "currency_choices": ForeignTenderProcess.CURRENCY_CHOICES,
                "selected_year": self.request.GET.get("year", ""),
                "selected_status": self.request.GET.get("status", ""),
                "selected_currency": self.request.GET.get("currency", ""),
                "search_query": self.request.GET.get("q", ""),
            }
        )
        return context


class ForeignTenderProcessCreateView(LoginRequiredMixin, GroupRequiredMixin, SuccessMessageMixin, CreateView):
    group_required = ["Supervisor", "Capturista"]
    model = ForeignTenderProcess
    form_class = ForeignTenderProcessForm
    template_name = "licitaciones/foreign_form.html"
    success_message = "Licitacion en el exterior creada correctamente."

    def form_valid(self, form):
        form.instance.created_by = self.request.user
        return super().form_valid(form)


class ForeignTenderProcessUpdateView(LoginRequiredMixin, GroupRequiredMixin, SuccessMessageMixin, UpdateView):
    group_required = ["Supervisor", "Capturista"]
    model = ForeignTenderProcess
    form_class = ForeignTenderProcessForm
    template_name = "licitaciones/foreign_form.html"
    success_message = "Licitacion en el exterior actualizada correctamente."


class ForeignTenderProcessDetailView(LoginRequiredMixin, DetailView):
    model = ForeignTenderProcess
    template_name = "licitaciones/foreign_detail.html"
    context_object_name = "process"

    def get_queryset(self):
        return ForeignTenderProcess.objects.select_related("created_by").prefetch_related(
            "requirements__unit",
            "purchase_orders",
            "updates__created_by",
        )


class ForeignTenderArchiveToggleView(LoginRequiredMixin, GroupRequiredMixin, View):
    group_required = ["Supervisor", "Capturista"]

    def post(self, request, pk):
        process = get_object_or_404(ForeignTenderProcess, pk=pk)
        process.is_active = not process.is_active
        process.save(update_fields=["is_active", "updated_at"])
        if process.is_active:
            messages.success(request, "Licitacion reactivada correctamente.")
            return redirect(process.get_absolute_url())
        messages.success(request, "Licitacion archivada correctamente.")
        return redirect("licitaciones:foreign_history")


class ForeignTenderRequirementCreateView(LoginRequiredMixin, SuccessMessageMixin, CreateView):
    model = ForeignTenderRequirement
    form_class = ForeignTenderRequirementForm
    template_name = "licitaciones/foreign_child_form.html"
    success_message = "Requerimiento agregado correctamente."

    def dispatch(self, request, *args, **kwargs):
        self.process = get_object_or_404(ForeignTenderProcess, pk=kwargs["process_pk"])
        return super().dispatch(request, *args, **kwargs)

    def form_valid(self, form):
        form.instance.process = self.process
        return super().form_valid(form)

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context.update({"process": self.process, "child_type": "requirement"})
        return context

    def get_success_url(self):
        return self.process.get_absolute_url()


class ForeignTenderRequirementUpdateView(LoginRequiredMixin, SuccessMessageMixin, UpdateView):
    model = ForeignTenderRequirement
    form_class = ForeignTenderRequirementForm
    template_name = "licitaciones/foreign_child_form.html"
    success_message = "Requerimiento actualizado correctamente."

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context.update({"process": self.object.process, "child_type": "requirement"})
        return context

    def get_success_url(self):
        return self.object.process.get_absolute_url()


class ForeignTenderUpdateCreateView(LoginRequiredMixin, SuccessMessageMixin, CreateView):
    model = ForeignTenderUpdate
    form_class = ForeignTenderUpdateForm
    template_name = "licitaciones/foreign_child_form.html"
    success_message = "Novedad documental registrada correctamente."

    def dispatch(self, request, *args, **kwargs):
        self.process = get_object_or_404(ForeignTenderProcess, pk=kwargs["process_pk"])
        return super().dispatch(request, *args, **kwargs)

    def form_valid(self, form):
        form.instance.process = self.process
        form.instance.created_by = self.request.user
        return super().form_valid(form)

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context.update({"process": self.process, "child_type": "update"})
        return context

    def get_success_url(self):
        return self.process.get_absolute_url()


class ForeignTenderUpdateUpdateView(LoginRequiredMixin, SuccessMessageMixin, UpdateView):
    model = ForeignTenderUpdate
    form_class = ForeignTenderUpdateForm
    template_name = "licitaciones/foreign_child_form.html"
    success_message = "Novedad documental actualizada correctamente."

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context.update({"process": self.object.process, "child_type": "update"})
        return context

    def get_success_url(self):
        return self.object.process.get_absolute_url()


class ForeignTenderPurchaseOrderCreateView(LoginRequiredMixin, SuccessMessageMixin, CreateView):
    model = ForeignTenderPurchaseOrder
    form_class = ForeignTenderPurchaseOrderForm
    template_name = "licitaciones/foreign_child_form.html"
    success_message = "Orden de compra agregada correctamente."

    def dispatch(self, request, *args, **kwargs):
        self.process = get_object_or_404(ForeignTenderProcess, pk=kwargs["process_pk"])
        return super().dispatch(request, *args, **kwargs)

    def form_valid(self, form):
        form.instance.process = self.process
        return super().form_valid(form)

    def get_form(self, form_class=None):
        form = super().get_form(form_class)
        if self.process.has_oca:
            form.fields.pop("saimb_number", None)
        return form

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context.update({"process": self.process, "child_type": "purchase_order"})
        return context

    def get_success_url(self):
        return self.process.get_absolute_url()


class ForeignTenderPurchaseOrderUpdateView(LoginRequiredMixin, SuccessMessageMixin, UpdateView):
    model = ForeignTenderPurchaseOrder
    form_class = ForeignTenderPurchaseOrderForm
    template_name = "licitaciones/foreign_child_form.html"
    success_message = "Orden de compra actualizada correctamente."

    def get_form(self, form_class=None):
        form = super().get_form(form_class)
        if self.object.process.has_oca:
            form.fields.pop("saimb_number", None)
        return form

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context.update({"process": self.object.process, "child_type": "purchase_order"})
        return context

    def get_success_url(self):
        return self.object.process.get_absolute_url()


class TenderStageManageView(LoginRequiredMixin, GroupRequiredMixin, DetailView):
    group_required = ["Supervisor", "Capturista"]
    model = TenderProcess
    template_name = "licitaciones/tender_stages.html"
    context_object_name = "process"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context["stages"] = self.object.stages.all().order_by("stage_number")
        return context

class TenderStageUpdateView(LoginRequiredMixin, GroupRequiredMixin, SuccessMessageMixin, UpdateView):
    group_required = ["Supervisor", "Capturista"]
    model = TenderStage
    form_class = TenderStageUpdateForm
    template_name = "licitaciones/tender_stage_form.html"
    success_message = "Etapa actualizada correctamente."

    def get_success_url(self):
        return reverse_lazy('licitaciones:tender_stages', kwargs={'pk': self.object.tender.pk})

@login_required
def export_national_tenders_csv(request):
    response = HttpResponse(content_type='text/csv; charset=utf-8-sig')
    response['Content-Disposition'] = 'attachment; filename="licitaciones_nacionales.csv"'

    writer = csv.writer(response, delimiter=';')
    writer.writerow(['Año', 'Unidad/Destino', 'Proceso', 'Expediente', 'Objeto', 'Clasificación', 'Estado', 'Monto Adjudicado (ARS)', 'IPP'])

    processes = TenderProcess.objects.select_related("unit").all().order_by("-year", "unit__name", "-opening_date", "process_number")

    for p in processes:
        writer.writerow([
            p.year,
            p.unit.name if p.unit else "SIN UNIDAD",
            p.process_number,
            p.expediente,
            p.name,
            p.get_classification_display() if p.classification else "-",
            p.get_status_display(),
            p.amount_ars if p.amount_ars is not None else "",
            p.ipp or "",
        ])

    return response


@login_required
def export_foreign_tenders_excel(request):
    queryset = ForeignTenderProcess.objects.filter(is_active=True).prefetch_related(
        "requirements__unit", "purchase_orders__provision_requests", "updates"
    ).order_by("-year", "process_number")

    year = _clean_int(request.GET.get("year"))
    status = request.GET.get("status", "")
    currency = request.GET.get("currency", "")
    query = request.GET.get("q", "").strip()

    if year:
        queryset = queryset.filter(year=year)
    if status:
        queryset = queryset.filter(status=status)
    if currency:
        queryset = queryset.filter(currency=currency)
    if query:
        queryset = queryset.filter(
            Q(process_number__icontains=query)
            | Q(expediente__icontains=query)
            | Q(requirements__requirement_number__icontains=query)
            | Q(requirements__description__icontains=query)
            | Q(purchase_orders__order_number__icontains=query)
            | Q(purchase_orders__supplier__icontains=query)
        ).distinct()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Procesos Exterior"

    header_fill = PatternFill(start_color="1F4B63", end_color="1F4B63", fill_type="solid")
    header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10)
    thin_border = Border(
        left=Side(style="thin", color="CCCCCC"),
        right=Side(style="thin", color="CCCCCC"),
        top=Side(style="thin", color="CCCCCC"),
        bottom=Side(style="thin", color="CCCCCC"),
    )
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    align_right = Alignment(horizontal="right", vertical="center", wrap_text=True)

    def apply_range_style(min_row, min_col, max_row, max_col, font=None, fill=None, border=None, alignment=None):
        for r in range(min_row, max_row + 1):
            for c in range(min_col, max_col + 1):
                cell = ws.cell(row=r, column=c)
                if font:
                    cell.font = font
                if fill:
                    cell.fill = fill
                if border:
                    cell.border = border
                if alignment:
                    cell.alignment = alignment

    # --- ENCABEZADOS (Fila 1 y Fila 2) ---
    ws.cell(row=1, column=1, value="DATOS INICIALES")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=4)

    ws.cell(row=2, column=1, value="LICITACION")
    ws.cell(row=2, column=2, value="OCA")
    ws.cell(row=2, column=3, value="REQ")
    ws.cell(row=2, column=4, value="OBJETO / EXPEDIENTE")

    headers_spanning = [
        (5, "MONTO\nREQUERIMIENTO"),
        (6, "MONTO DICTAMEN\nDE EVALUACION"),
        (7, "TALLER"),
        (8, "AERONAVE"),
        (9, "DESCRIPCION"),
        (10, "ESTADO DE LA\nCONTRATACION"),
        (11, "MONTO\nASIGNADO"),
        (12, "GFH DE\nASIGNACION"),
        (13, "INCOTERM"),
        (14, "NRO. OC"),
        (15, "ORDEN DE\nCOMPRA"),
        (16, "PROVEEDOR"),
        (17, "MONTO OC\nCOMPROMETIDO"),
        (18, "FECHA DE\nEMISION"),
        (19, "FECHA VTO OC"),
        (20, "SOLICITUD DE\nPROVISION (SP)"),
        (21, "MONTO SP"),
        (22, "FECHA EMISION\n(SP)"),
        (23, "FECHA VTO\n(SP)"),
        (24, "MONTO\nREMANENTE"),
        (25, "SAIMB NRO."),
        (26, "RECIBIDO"),
        (27, "ULTIMO ESTADO\nPOR GDE/GFH"),
    ]

    for col_idx, text in headers_spanning:
        ws.cell(row=1, column=col_idx, value=text)
        ws.merge_cells(start_row=1, start_column=col_idx, end_row=2, end_column=col_idx)

    apply_range_style(
        min_row=1,
        min_col=1,
        max_row=2,
        max_col=27,
        font=header_font,
        fill=header_fill,
        border=thin_border,
        alignment=align_center,
    )
    ws.row_dimensions[1].height = 26
    ws.row_dimensions[2].height = 24

    # --- DATOS ---
    row_num = 3
    for process in queryset:
        pos = list(process.purchase_orders.all())
        reqs = list(process.requirements.all())
        n_reqs = len(reqs)

        oc_numbers = "\n".join(po.order_number for po in pos) if pos else "-"
        oc_types = "\n".join("OCA" if po.order_type == "OCA" else "OC" for po in pos) if pos else "-"
        oc_suppliers = "\n".join(po.supplier or "-" for po in pos) if pos else "-"
        oc_amounts = "\n".join(f"{process.currency_symbol} {po.amount:,.2f}" for po in pos if po.amount is not None) if pos else "-"
        oc_issues = "\n".join(po.issue_date.strftime("%d/%m/%Y") for po in pos if po.issue_date) if pos else "-"
        oc_exps = "\n".join(po.expiration_date.strftime("%d/%m/%Y") for po in pos if po.expiration_date) if pos else "-"

        all_sps = [sp for po in pos for sp in po.provision_requests.all()]
        sp_numbers = "\n".join(sp.sp_number for sp in all_sps) if all_sps else "-"
        sp_amounts = "\n".join(f"{process.currency_symbol} {sp.amount:,.2f}" for sp in all_sps if sp.amount is not None) if all_sps else "-"
        sp_issues = "\n".join(sp.issue_date.strftime("%d/%m/%Y") for sp in all_sps if sp.issue_date) if all_sps else "-"
        sp_exps = "\n".join(sp.expiration_date.strftime("%d/%m/%Y") for sp in all_sps if sp.expiration_date) if all_sps else "-"

        if process.has_oca:
            saimb = "\n".join(sp.saimb_number for sp in all_sps if sp.saimb_number) if all_sps else "-"
            recibido = "\n".join("SI" if sp.received else ("NO" if sp.received is False else "-") for sp in all_sps) if all_sps else "-"
        else:
            saimb = "\n".join(po.saimb_number for po in pos if po.saimb_number) if pos else "-"
            recibido = "SI" if process.received else ("NO" if process.received is False else "-")

        if process.latest_update:
            org = f"{process.latest_update.organization} - " if process.latest_update.organization else ""
            date_str = f" ({process.latest_update.event_date.strftime('%d/%m/%Y')})" if process.latest_update.event_date else ""
            ultimo_estado = f"{org}{process.latest_update.description}{date_str}"
        else:
            ultimo_estado = process.notes or "-"

        has_oca_str = "SI" if process.has_oca else ("NO" if process.has_oca is False else "-")
        eval_amount_str = f"{process.currency_symbol} {process.evaluation_amount:,.2f}" if process.evaluation_amount is not None else "-"
        awarded_amount_str = f"{process.currency_symbol} {process.awarded_amount:,.2f}" if process.awarded_amount is not None else "-"
        remaining_amount_str = f"{process.currency_symbol} {process.remaining_amount:,.2f}" if process.remaining_amount is not None else "-"
        licitacion_str = f"{process.process_number}\n{process.year}"

        if n_reqs > 0:
            start_row = row_num
            end_row = row_num + n_reqs - 1

            for idx, req in enumerate(reqs):
                r = start_row + idx
                req_amount_str = f"{process.currency_symbol} {req.requested_amount:,.2f}" if req.requested_amount is not None else "-"
                ws.cell(row=r, column=3, value=req.requirement_number)
                ws.cell(row=r, column=5, value=req_amount_str)
                ws.cell(row=r, column=7, value=req.workshop_label or "-")
                ws.cell(row=r, column=8, value=req.aircraft or "-")
                ws.cell(row=r, column=9, value=req.description or "-")

            # Valores del proceso en start_row
            ws.cell(row=start_row, column=1, value=licitacion_str)
            ws.cell(row=start_row, column=2, value=has_oca_str)
            ws.cell(row=start_row, column=4, value=process.expediente or "-")
            ws.cell(row=start_row, column=6, value=eval_amount_str)
            ws.cell(row=start_row, column=10, value=process.get_status_display())
            ws.cell(row=start_row, column=11, value=awarded_amount_str)
            ws.cell(row=start_row, column=12, value=process.allocation_gfh or "-")
            ws.cell(row=start_row, column=13, value=process.incoterm or "-")
            ws.cell(row=start_row, column=14, value=oc_numbers)
            ws.cell(row=start_row, column=15, value=oc_types)
            ws.cell(row=start_row, column=16, value=oc_suppliers)
            ws.cell(row=start_row, column=17, value=oc_amounts)
            ws.cell(row=start_row, column=18, value=oc_issues)
            ws.cell(row=start_row, column=19, value=oc_exps)
            ws.cell(row=start_row, column=20, value=sp_numbers)
            ws.cell(row=start_row, column=21, value=sp_amounts)
            ws.cell(row=start_row, column=22, value=sp_issues)
            ws.cell(row=start_row, column=23, value=sp_exps)
            ws.cell(row=start_row, column=24, value=remaining_amount_str)
            ws.cell(row=start_row, column=25, value=saimb)
            ws.cell(row=start_row, column=26, value=recibido)
            ws.cell(row=start_row, column=27, value=ultimo_estado)

            if n_reqs > 1:
                cols_to_merge = [1, 2, 4, 6, 10, 11, 12, 13, 14, 15, 16, 17, 18, 19, 20, 21, 22, 23, 24, 25, 26, 27]
                for col_c in cols_to_merge:
                    ws.merge_cells(start_row=start_row, start_column=col_c, end_row=end_row, end_column=col_c)

            for r in range(start_row, end_row + 1):
                for col_c in range(1, 28):
                    c = ws.cell(row=r, column=col_c)
                    c.font = data_font
                    c.border = thin_border
                    if col_c in (4, 9, 27):
                        c.alignment = align_left
                    elif col_c in (5, 6, 11, 17, 21, 24):
                        c.alignment = align_right
                    else:
                        c.alignment = align_center

            row_num = end_row + 1
        else:
            ws.cell(row=row_num, column=1, value=licitacion_str)
            ws.cell(row=row_num, column=2, value=has_oca_str)
            ws.cell(row=row_num, column=3, value="-")
            ws.cell(row=row_num, column=4, value=process.expediente or "-")
            ws.cell(row=row_num, column=5, value="-")
            ws.cell(row=row_num, column=6, value=eval_amount_str)
            ws.cell(row=row_num, column=7, value="-")
            ws.cell(row=row_num, column=8, value="-")
            ws.cell(row=row_num, column=9, value="-")
            ws.cell(row=row_num, column=10, value=process.get_status_display())
            ws.cell(row=row_num, column=11, value=awarded_amount_str)
            ws.cell(row=row_num, column=12, value=process.allocation_gfh or "-")
            ws.cell(row=row_num, column=13, value=process.incoterm or "-")
            ws.cell(row=row_num, column=14, value=oc_numbers)
            ws.cell(row=row_num, column=15, value=oc_types)
            ws.cell(row=row_num, column=16, value=oc_suppliers)
            ws.cell(row=row_num, column=17, value=oc_amounts)
            ws.cell(row=row_num, column=18, value=oc_issues)
            ws.cell(row=row_num, column=19, value=oc_exps)
            ws.cell(row=row_num, column=20, value=sp_numbers)
            ws.cell(row=row_num, column=21, value=sp_amounts)
            ws.cell(row=row_num, column=22, value=sp_issues)
            ws.cell(row=row_num, column=23, value=sp_exps)
            ws.cell(row=row_num, column=24, value=remaining_amount_str)
            ws.cell(row=row_num, column=25, value=saimb)
            ws.cell(row=row_num, column=26, value=recibido)
            ws.cell(row=row_num, column=27, value=ultimo_estado)

            for col_c in range(1, 28):
                c = ws.cell(row=row_num, column=col_c)
                c.font = data_font
                c.border = thin_border
                if col_c in (4, 9, 27):
                    c.alignment = align_left
                elif col_c in (5, 6, 11, 17, 21, 24):
                    c.alignment = align_right
                else:
                    c.alignment = align_center
            row_num += 1

    # Ajuste dinámico del ancho de columnas
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val = str(cell.value or "")
            for l in val.split("\n"):
                if len(l) > max_len:
                    max_len = len(l)
        ws.column_dimensions[col_letter].width = min(max(max_len + 4, 14), 50)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = 'attachment; filename="licitaciones_exteriores.xlsx"'
    return response

@login_required
def mark_notification_read(request, pk):
    from .models import Notification
    try:
        notification = Notification.objects.get(pk=pk, user=request.user)
        notification.is_read = True
        notification.save()
        if notification.link:
            return redirect(notification.link)
    except Notification.DoesNotExist:
        pass
    return redirect('licitaciones:type_selection')

class ForeignProvisionRequestCreateView(LoginRequiredMixin, SuccessMessageMixin, CreateView):
    model = ForeignProvisionRequest
    form_class = ForeignProvisionRequestForm
    template_name = 'licitaciones/foreign_child_form.html'
    success_message = 'Solicitud de provisión agregada correctamente.'

    def dispatch(self, request, *args, **kwargs):
        self.purchase_order = get_object_or_404(ForeignTenderPurchaseOrder, pk=kwargs['order_pk'])
        if not self.purchase_order.process.has_oca:
            messages.error(request, "Las Órdenes de Compra cerradas no admiten Solicitudes de Provisión.")
            return redirect(self.purchase_order.process.get_absolute_url())
        return super().dispatch(request, *args, **kwargs)

    def form_valid(self, form):
        form.instance.purchase_order = self.purchase_order
        return super().form_valid(form)

    def get_form(self, form_class=None):
        form = super().get_form(form_class)
        if not self.purchase_order.process.has_oca:
            form.fields.pop("saimb_number", None)
        return form

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context.update({'process': self.purchase_order.process, 'child_type': 'provision_request'})
        return context

    def get_success_url(self):
        return self.purchase_order.process.get_absolute_url()


class ForeignProvisionRequestUpdateView(LoginRequiredMixin, SuccessMessageMixin, UpdateView):
    model = ForeignProvisionRequest
    form_class = ForeignProvisionRequestForm
    template_name = 'licitaciones/foreign_child_form.html'
    success_message = 'Solicitud de provisión actualizada correctamente.'

    def get_form(self, form_class=None):
        form = super().get_form(form_class)
        if not self.object.purchase_order.process.has_oca:
            form.fields.pop("saimb_number", None)
        return form

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context.update({'process': self.object.purchase_order.process, 'child_type': 'provision_request'})
        return context

    def get_success_url(self):
        return self.object.purchase_order.process.get_absolute_url()


class ForeignTenderDeleteView(LoginRequiredMixin, View):
    def dispatch(self, request, *args, **kwargs):
        return self._dispatch(request, *args, **kwargs)

    def _get_target_object(self, model_type, pk):
        mapping = {
            "process": (ForeignTenderProcess, "Licitación"),
            "requirement": (ForeignTenderRequirement, "Requerimiento"),
            "purchase_order": (ForeignTenderPurchaseOrder, "Orden de Compra"),
            "provision_request": (ForeignProvisionRequest, "Solicitud de Provisión"),
            "update": (ForeignTenderUpdate, "Novedad Documental"),
        }
        if model_type not in mapping:
            return None, None
        model_cls, label = mapping[model_type]
        obj = get_object_or_404(model_cls, pk=pk)
        return obj, label

    def _get_redirect_url(self, obj, model_type):
        if model_type == "process":
            return reverse_lazy("licitaciones:foreign_list")
        elif model_type in ("requirement", "purchase_order", "update"):
            return obj.process.get_absolute_url()
        elif model_type == "provision_request":
            return obj.purchase_order.process.get_absolute_url()
        return reverse_lazy("licitaciones:foreign_list")

    def _dispatch(self, request, *args, **kwargs):
        model_type = kwargs.get("model_type")
        pk = kwargs.get("pk")
        obj, label = self._get_target_object(model_type, pk)
        if not obj:
            messages.error(request, "Tipo de objeto no válido.")
            return redirect("licitaciones:foreign_list")

        redirect_url = self._get_redirect_url(obj, model_type)

        from core.models import UserSystemPIN
        if not UserSystemPIN.objects.filter(user=request.user, system_code="procurement_delete").exists():
            messages.error(request, "No tiene configurado el PIN de Borrado de Compras.")
            return redirect(redirect_url)

        if request.method == "GET":
            return render(
                request,
                "licitaciones/foreign_confirm_delete.html",
                {
                    "object": obj,
                    "label": label,
                    "model_type": model_type,
                    "redirect_url": redirect_url,
                },
            )
        elif request.method == "POST":
            confirmation = request.POST.get("confirmation", "").strip()
            pin = request.POST.get("pin", "")

            if confirmation.upper() != "BORRAR":
                messages.error(request, 'Debe escribir "BORRAR" en mayúsculas para confirmar.')
                return render(
                    request,
                    "licitaciones/foreign_confirm_delete.html",
                    {
                        "object": obj,
                        "label": label,
                        "model_type": model_type,
                        "redirect_url": redirect_url,
                    },
                )

            from django.contrib.auth.hashers import check_password
            try:
                access = UserSystemPIN.objects.get(user=request.user, system_code="procurement_delete")
                pin_match = check_password(pin, access.pin_hash)
            except UserSystemPIN.DoesNotExist:
                pin_match = False

            if not pin_match:
                messages.error(request, "PIN incorrecto. No se realizó la eliminación.")
                return render(
                    request,
                    "licitaciones/foreign_confirm_delete.html",
                    {
                        "object": obj,
                        "label": label,
                        "model_type": model_type,
                        "redirect_url": redirect_url,
                    },
                )

            try:
                obj_name = str(obj)
                obj.delete()
                messages.success(request, f"{label} '{obj_name}' eliminado/a correctamente.")
            except Exception as e:
                messages.error(request, f"No se pudo eliminar el registro: {e}")
            return redirect(redirect_url)

